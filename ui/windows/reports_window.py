from __future__ import annotations
from PySide2.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout,
                                QPushButton, QLabel, QFrame, QFileDialog,
                                QMessageBox, QTableWidget, QTableWidgetItem,
                                QHeaderView, QTabWidget)
from PySide2.QtCore import Qt
from PySide2.QtGui import QColor
from app.database.connection import get_session
from app.models.sale import Sale
from app.repositories.product_repository import ProductRepository
from app.repositories.customer_repository import CustomerRepository
from app.repositories.supplier_repository import SupplierRepository
from app.repositories.sale_repository import SaleRepository
from app.services.billing_service import BillingService
from app.config.settings import REPORTS_DIR
from app.design.icons import icon
from ui.widgets.table_action_button import TableActionButton
import os


class ReportsWindow(QWidget):
    def __init__(self, user_data: dict = None):
        super().__init__()
        self.user_data = user_data or {}
        self._setup_ui()
        self.load_stats()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 16, 24, 24)
        layout.setSpacing(16)

        cards_row = QHBoxLayout()
        cards_row.setSpacing(12)
        self.card_products  = self._card("Total Products",  "0", "#89b4fa", "products")
        self.card_customers = self._card("Customers",        "0", "#a6e3a1", "customers")
        self.card_suppliers = self._card("Suppliers",        "0", "#cba6f7", "suppliers")
        self.card_sales     = self._card("Total Invoices",   "0", "#fab387", "receipt")
        self.card_revenue   = self._card("Total Revenue",   "₹0", "#f9e2af", "payment")
        self.card_lowstock  = self._card("Low Stock",        "0", "#f38ba8", "warning")
        for c in [self.card_products, self.card_customers, self.card_suppliers,
                  self.card_sales, self.card_revenue, self.card_lowstock]:
            cards_row.addWidget(c["frame"])

        export_row = QHBoxLayout()
        self.btn_excel_products  = self._export_btn("Products → Excel",  "#89b4fa", "export")
        self.btn_excel_sales     = self._export_btn("Sales → Excel",     "#a6e3a1", "export")
        self.btn_excel_customers = self._export_btn("Customers → Excel", "#cba6f7", "export")
        self.btn_pdf_summary     = self._export_btn("Summary → PDF",     "#fab387", "receipt")
        self.btn_refresh         = QPushButton(" Refresh")
        self.btn_refresh.setIcon(icon("refresh"))

        self.btn_excel_products.clicked.connect(self._export_products_excel)
        self.btn_excel_sales.clicked.connect(self._export_sales_excel)
        self.btn_excel_customers.clicked.connect(self._export_customers_excel)
        self.btn_pdf_summary.clicked.connect(self._export_pdf_summary)
        self.btn_refresh.clicked.connect(self.load_stats)

        export_row.addWidget(self.btn_excel_products)
        export_row.addWidget(self.btn_excel_sales)
        export_row.addWidget(self.btn_excel_customers)
        export_row.addWidget(self.btn_pdf_summary)
        export_row.addStretch()
        export_row.addWidget(self.btn_refresh)

        tabs = QTabWidget()
        tabs.addTab(self._build_product_table(),  icon("products"), "Products")
        tabs.addTab(self._build_sales_table(),    icon("sales"), "Sales")
        tabs.addTab(self._build_customer_table(), icon("customers"), "Customers")
        tabs.setStyleSheet("""
            QTabBar::tab {
                font-size: 16px;
                font-weight: bold;
                padding: 10px 24px;
            }
        """)

        layout.addLayout(cards_row)
        layout.addLayout(export_row)
        layout.addWidget(tabs)

    def _card(self, label, value, color, icon_name=None) -> dict:
        frame = QFrame()
        frame.setObjectName("StatCard")
        frame.setMinimumHeight(90)
        inner = QVBoxLayout(frame)
        inner.setContentsMargins(14, 10, 14, 10)
        inner.setSpacing(4)
        val_lbl = QLabel(value)
        val_lbl.setStyleSheet(f"font-size: 22px; font-weight: bold; color: {color};")

        label_row = QHBoxLayout()
        label_row.setSpacing(6)
        if icon_name:
            icon_lbl = QLabel()
            icon_lbl.setPixmap(icon(icon_name, color=color).pixmap(14, 14))
            label_row.addWidget(icon_lbl)
        lbl = QLabel(label)
        lbl.setObjectName("StatLabel")
        label_row.addWidget(lbl)
        label_row.addStretch()

        inner.addWidget(val_lbl)
        inner.addLayout(label_row)
        return {"frame": frame, "val": val_lbl}

    def _export_btn(self, label, color, icon_name=None) -> QPushButton:
        btn = QPushButton(label)
        btn.setObjectName("SmallButton")
        if icon_name:
            btn.setIcon(icon(icon_name, color=color))
        return btn

    def _build_product_table(self) -> QWidget:
        w = QWidget()
        l = QVBoxLayout(w)
        self.product_table = QTableWidget()
        self.product_table.setColumnCount(6)
        self.product_table.setHorizontalHeaderLabels(["Name", "SKU", "Size", "Price", "Stock", "Status"])
        self.product_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        l.addWidget(self.product_table)
        return w

    def _build_sales_table(self) -> QWidget:
        w = QWidget()
        l = QVBoxLayout(w)
        self.sales_table = QTableWidget()
        self.sales_table.setColumnCount(7)
        self.sales_table.setHorizontalHeaderLabels(
            ["Date", "Items (Qty)", "Customer", "Subtotal", "GST", "Total Amount", "Action"]
        )
        self.sales_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.sales_table.verticalHeader().setDefaultSectionSize(50)
        l.addWidget(self.sales_table)
        return w

    def _build_customer_table(self) -> QWidget:
        w = QWidget()
        l = QVBoxLayout(w)
        self.customer_table = QTableWidget()
        self.customer_table.setColumnCount(4)
        self.customer_table.setHorizontalHeaderLabels(["Name", "Email", "Phone", "Joined"])
        self.customer_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        l.addWidget(self.customer_table)
        return w

    def load_stats(self):
        session = get_session()
        try:
            products  = ProductRepository(session).get_all()
            customers = CustomerRepository(session).get_all()
            suppliers = SupplierRepository(session).get_all()
            sales     = SaleRepository(session).get_all()
            revenue   = sum(
                s.total_amount for s in sales
                if not (hasattr(s, "status") and s.status == "Cancelled")
            )
            low_stock_count = sum(1 for p in products if 0 < p.stock_qty <= 10)

            self.card_products["val"].setText(str(len(products)))
            self.card_customers["val"].setText(str(len(customers)))
            self.card_suppliers["val"].setText(str(len(suppliers)))
            self.card_sales["val"].setText(str(len(sales)))
            self.card_revenue["val"].setText(f"₹{revenue:,.0f}")
            self.card_lowstock["val"].setText(str(low_stock_count))

            self.product_table.setRowCount(len(products))
            for row, p in enumerate(products):
                self.product_table.setItem(row, 0, QTableWidgetItem(str(p.name)))
                self.product_table.setItem(row, 1, QTableWidgetItem(str(p.sku)))
                self.product_table.setItem(row, 2, QTableWidgetItem(str(getattr(p, "size", "N/A"))))
                self.product_table.setItem(row, 3, QTableWidgetItem(f"₹{p.price:.2f}"))
                self.product_table.setItem(row, 4, QTableWidgetItem(str(p.stock_qty)))
                self.product_table.setItem(
                    row, 5,
                    QTableWidgetItem(p.status.value if hasattr(p.status, "value") else str(p.status))
                )

            self.customer_table.setRowCount(len(customers))
            for row, c in enumerate(customers):
                self.customer_table.setItem(row, 0, QTableWidgetItem(str(c.name)))
                self.customer_table.setItem(row, 1, QTableWidgetItem(str(c.email or "")))
                self.customer_table.setItem(row, 2, QTableWidgetItem(str(c.phone or "")))
                joined = (
                    c.created_at.strftime("%d-%m-%Y")
                    if hasattr(c, "created_at") and c.created_at
                    else "N/A"
                )
                self.customer_table.setItem(row, 3, QTableWidgetItem(joined))

            self.sales_table.setRowCount(len(sales))
            for row, s in enumerate(sales):
                items_list = [
                    f"{item.product_name} (x{item.quantity})"
                    for item in getattr(s, "items", [])
                ]
                self.sales_table.setItem(row, 0, QTableWidgetItem(s.date_created.strftime("%d-%m-%Y")))
                self.sales_table.setItem(row, 1, QTableWidgetItem(", ".join(items_list)))
                self.sales_table.setItem(row, 2, QTableWidgetItem(s.customer_name or "Walk-in"))
                self.sales_table.setItem(row, 3, QTableWidgetItem(f"₹{s.subtotal:.2f}"))
                self.sales_table.setItem(row, 4, QTableWidgetItem(f"₹{s.gst_amount:.2f}"))
                self.sales_table.setItem(row, 5, QTableWidgetItem(f"₹{s.total_amount:.2f}"))

                action_widget = QWidget()
                action_layout = QHBoxLayout(action_widget)
                action_layout.setContentsMargins(2, 2, 2, 2)
                action_layout.setSpacing(4)

                reprint_btn = TableActionButton("reprint", "Reprint invoice",
                                                 on_click=lambda checked=False, sid=s.id: self._reprint_invoice(sid))
                action_layout.addWidget(reprint_btn)

                if hasattr(s, "status") and s.status == "Cancelled":
                    cancel_lbl = QLabel("Cancelled")
                    cancel_lbl.setStyleSheet("color: #f38ba8; font-weight: bold;")
                    cancel_lbl.setAlignment(Qt.AlignCenter)
                    action_layout.addWidget(cancel_lbl)
                else:
                    cancel_btn = QPushButton(" Cancel Bill")
                    cancel_btn.setStyleSheet("background-color: #a33b3b; color: white;")
                    cancel_btn.clicked.connect(lambda checked=False, sid=s.id: self._cancel_invoice_ui(sid))
                    action_layout.addWidget(cancel_btn)

                self.sales_table.setCellWidget(row, 6, action_widget)

        finally:
            session.close()

    def _cancel_invoice_ui(self, sale_id):
        from app.services.billing_service import BillingService

        if (
            QMessageBox.warning(
                self, "Cancel",
                "Permanently delete and restore stock?",
                QMessageBox.Yes | QMessageBox.No,
            )
            == QMessageBox.Yes
        ):
            session = get_session()
            try:
                BillingService.cancel_invoice(session, sale_id)
                self.load_stats()
            finally:
                session.close()

    def _reprint_invoice(self, sale_id):
        from app.utils.pdf_generator import generate_invoice_pdf

        session = get_session()
        try:
            sale = SaleRepository(session).get_by_id(sale_id)
            if not sale:
                QMessageBox.warning(self, "Not Found", "This invoice could not be found.")
                return

            pdf_path = generate_invoice_pdf(
                sale,
                company_name=self.user_data.get("company_name", "Dindigul School Point"),
                company_address=self.user_data.get("company_address", ""),
                company_gst=self.user_data.get("company_gst", ""),
                company_phone=self.user_data.get("company_phone", ""),
                company_email=self.user_data.get("email", ""),
                company_logo_path=self.user_data.get("company_logo_path", ""),
            )

            printed = BillingService.print_pdf(
                pdf_path,
                preferred_printer=self.user_data.get("preferred_printer", ""),
            )
            if printed:
                QMessageBox.information(self, "Reprinted", f"Invoice #{sale_id} sent to printer.")
            else:
                QMessageBox.information(
                    self, "Saved",
                    f"Invoice #{sale_id} saved as PDF (no printer found):\n{pdf_path}"
                )
        finally:
            session.close()

    def _export_products_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            session = get_session()
            products = ProductRepository(session).get_all()
            session.close()

            path, _ = QFileDialog.getSaveFileName(
                self, "Save Products Excel",
                os.path.join(REPORTS_DIR, "products.xlsx"),
                "Excel Files (*.xlsx)",
            )
            if not path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Products"

            headers = ["ID", "Name", "SKU", "Size", "Price", "Stock", "Status"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
                cell.alignment = Alignment(horizontal="center")

            for row_idx, p in enumerate(products, 2):
                ws.cell(row=row_idx, column=1, value=p.id)
                ws.cell(row=row_idx, column=2, value=p.name)
                ws.cell(row=row_idx, column=3, value=p.sku)
                ws.cell(row=row_idx, column=4, value=getattr(p, "size", "N/A"))
                ws.cell(row=row_idx, column=5, value=p.price)
                ws.cell(row=row_idx, column=6, value=p.stock_qty)
                ws.cell(
                    row=row_idx, column=7,
                    value=p.status.value if hasattr(p.status, "value") else str(p.status),
                )

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 18

            wb.save(path)
            QMessageBox.information(self, "Exported", f"Products exported to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Failed", str(e))

    def _export_sales_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from sqlalchemy.orm import joinedload

            session = get_session()
            # Eagerly load items while session is open
            sales = session.query(Sale).options(joinedload(Sale.items)).all()
            session.close()

            path, _ = QFileDialog.getSaveFileName(
                self, "Save Sales Excel",
                os.path.join(REPORTS_DIR, "sales.xlsx"),
                "Excel Files (*.xlsx)",
            )
            if not path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sales"

            headers = [
                "Invoice ID", "Date", "Items", "Customer",
                "Subtotal (INR)", "GST (INR)", "Total Amount (INR)", "Status",
            ]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="16A34A")
                cell.alignment = Alignment(horizontal="center")

            for row_idx, s in enumerate(sales, 2):
                items_list = [item.product_name for item in getattr(s, "items", [])]
                items_str = ", ".join(items_list) if items_list else "Unknown Items"
                status_str = s.status if hasattr(s, "status") else "Active"

                ws.cell(row=row_idx, column=1, value=s.id)
                ws.cell(
                    row=row_idx, column=2,
                    value=(
                        s.date_created.strftime("%d-%m-%Y %H:%M")
                        if hasattr(s, "date_created") and s.date_created
                        else ""
                    ),
                )
                ws.cell(row=row_idx, column=3, value=items_str)
                ws.cell(row=row_idx, column=4, value=s.customer_name or "Walk-in")
                ws.cell(row=row_idx, column=5, value=s.subtotal)
                ws.cell(row=row_idx, column=6, value=s.gst_amount)
                ws.cell(row=row_idx, column=7, value=s.total_amount)
                ws.cell(row=row_idx, column=8, value=status_str)

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 20

            wb.save(path)
            QMessageBox.information(self, "Exported", f"Sales exported to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Failed", str(e))

    def _export_customers_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            session = get_session()
            customers = CustomerRepository(session).get_all()
            session.close()

            path, _ = QFileDialog.getSaveFileName(
                self, "Save Customers Excel",
                os.path.join(REPORTS_DIR, "customers.xlsx"),
                "Excel Files (*.xlsx)",
            )
            if not path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Customers"

            headers = ["ID", "Name", "Email", "Phone", "Address", "Joined"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="7C3AED")
                cell.alignment = Alignment(horizontal="center")

            for row_idx, c in enumerate(customers, 2):
                ws.cell(row=row_idx, column=1, value=c.id)
                ws.cell(row=row_idx, column=2, value=c.name)
                ws.cell(row=row_idx, column=3, value=c.email or "")
                ws.cell(row=row_idx, column=4, value=c.phone or "")
                ws.cell(row=row_idx, column=5, value=c.address or "")
                ws.cell(
                    row=row_idx, column=6,
                    value=c.created_at.strftime("%d-%m-%Y") if c.created_at else "",
                )

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 20

            wb.save(path)
            QMessageBox.information(self, "Exported", f"Customers exported to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Failed", str(e))

    def _export_pdf_summary(self):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from datetime import datetime

            session = get_session()
            products  = ProductRepository(session).get_all()
            customers = CustomerRepository(session).get_all()
            sales     = SaleRepository(session).get_all()
            session.close()

            path, _ = QFileDialog.getSaveFileName(
                self, "Save PDF Summary",
                os.path.join(REPORTS_DIR, "summary.pdf"),
                "PDF Files (*.pdf)",
            )
            if not path:
                return

            doc    = SimpleDocTemplate(path, pagesize=A4)
            styles = getSampleStyleSheet()
            story  = []

            story.append(Paragraph("Stock Manager — Summary Report", styles["Title"]))
            story.append(
                Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}", styles["Normal"])
            )
            story.append(Spacer(1, 20))

            revenue = sum(
                s.total_amount for s in sales
                if not (hasattr(s, "status") and s.status == "Cancelled")
            )
            summary_data = [
                ["Metric", "Value"],
                ["Total Products",  str(len(products))],
                ["Total Customers", str(len(customers))],
                ["Total Invoices",  str(len(sales))],
                ["Total Revenue",   f"₹ {revenue:,.2f}"],
                ["Low Stock Items", str(sum(1 for p in products if 0 < p.stock_qty <= 10))],
                ["Out of Stock",    str(sum(1 for p in products if p.stock_qty == 0))],
            ]
            t = Table(summary_data, colWidths=[250, 200])
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
                        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                        ("PADDING", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 24))

            story.append(Paragraph("Products List", styles["Heading2"]))
            story.append(Spacer(1, 8))
            prod_data = [["Name", "SKU", "Size", "Price", "Stock", "Status"]]
            for p in products[:30]:
                prod_data.append(
                    [
                        p.name[:30],
                        p.sku,
                        getattr(p, "size", "N/A"),
                        f"₹ {p.price:.2f}",
                        str(p.stock_qty),
                        p.status.value if hasattr(p.status, "value") else str(p.status),
                    ]
                )
            pt = Table(prod_data, colWidths=[160, 70, 60, 70, 50, 70])
            pt.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
                        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F1F5F9"), colors.white]),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("PADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(pt)
            doc.build(story)
            QMessageBox.information(self, "Exported", f"PDF saved to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Failed", str(e))

    def refresh(self):
        self.load_stats()